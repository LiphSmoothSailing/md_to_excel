# -*- coding: utf-8 -*-
"""MD → 設計書 Excel 生成器（基于 design_template.xlsx 版式）。

两种输入模式：
  1) --input xxx.md     解析 Markdown（保留原文语言）渲染
  2) --blocks xxx.json  直接渲染 blocks JSON（适合 agent 预先翻译/组织内容）

blocks JSON 结构（列表）：
  ["chapter", "标题文本"]                      章节标题（深灰底白字）
  ["section", "标题文本"]                      子节标题（浅灰底）
  ["table", ["列1","列2"], [["a","b"], ...]]   表格（1~3 列；>3 列自动并入第 3 列）
  ["note", "文本"]                             正文段落（可含 \\n 多行）
  ["code", "文本"]                             代码/JSON 示例（浅灰底 Consolas）
  ["mermaid", "mermaid代码"]                   流程图（渲染 PNG 嵌入；失败降级为代码文本并输出告警）

表格单元格中：`<br>`（大小写不敏感，可写 `<br/>`）转为单元格内换行；`\\|` 还原为 `|`。

模板版式铁律（不可违反）：
  - 保留模板前 6 行（值/样式/合并/行高）
  - 列宽与模板完全一致（A=3.25、AF=4.5、AG=3.25，其余默认 0.8cm≈4.3 字符），绝不改列宽
  - 长内容一律用合并单元格承载

目录约定（项目内）：
  - 模板：Skill 内 assets/design_template.xlsx（唯一权威，只留这一份）
  - 临时文件：./temp/（转换中产生的 mermaid 图片等）
  - 输出：--output 未指定时默认 ./output/<源文件名>.xlsx
"""
import argparse
import json
import math
import os
import re
import shutil
import struct
import subprocess
import sys
import tempfile
import unicodedata
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MMDC = shutil.which('mmdc')

# 模板唯一位置：本 Skill 的 assets/design_template.xlsx（<skill>/assets/）
ASSET_TPL = Path(__file__).resolve().parent.parent / 'assets' / 'design_template.xlsx'

FONT_JA = 'Meiryo'
FONT_MONO = 'Consolas'
LINE_THIN = Side(style='thin', color='FF000000')

# 内容区列段（1-based 列号），基于模板 B..BC 列宽体系
SEG3 = [(2, 5), (6, 29), (30, 55)]   # B:E | F:AC | AD:BC
SEG2 = [(2, 8), (9, 55)]             # B:H | I:BC
SEG1 = [(2, 55)]                     # B:BC
FULL = (2, 55)

DEFAULT_SHEET = '設計書'
DEFAULT_W = 4.32  # 模板默认列宽（0.8cm ≈ 4.3 字符）

# 行高估算（经验值，与既有版式保持一致）
EST_AVAIL_DIV = 2          # 列宽字符数 → 可容纳显示宽度的折算除数
NOTE_LINE_PT, NOTE_PAD = 14, 8
CODE_LINE_PT, CODE_PAD = 13, 8
TABLE_LINE_PT, TABLE_PAD = 14, 6

MERMAID_WIDTH = 600
MERMAID_MAX_H = 520        # Excel 行高上限约 409.5pt ≈ 546px，保守取 520
MMDC_TIMEOUT = 60          # 秒；浏览器卡死时避免无限等待

BLOCK_KINDS = ('chapter', 'section', 'table', 'note', 'code', 'mermaid')
BR_RE = re.compile(r'<br\s*/?>', re.IGNORECASE)


def disp_width(s):
    return sum(2 if unicodedata.east_asian_width(ch) in ('F', 'W') else 1
               for ch in str(s or ''))


def col_width(ws, c):
    """取列实际宽度（显式列宽优先，否则默认列宽）"""
    w = ws.column_dimensions[get_column_letter(c)].width
    return float(w) if w else float(ws.sheet_format.defaultColWidth)


def seg_width(ws, c1, c2):
    """合并段实际总列宽（字符）"""
    return sum(col_width(ws, c) for c in range(c1, c2 + 1))


def est_lines(text, c1, c2, ws=None):
    text = str(text or '')
    if not text:
        return 1
    wsum = seg_width(ws, c1, c2) if ws is not None else DEFAULT_W * (c2 - c1 + 1)
    avail = max(wsum / EST_AVAIL_DIV - 2, 8)
    return sum(max(1, math.ceil(disp_width(ln) / avail)) for ln in text.split('\n'))


def dynamic_segments(ws, n, headers, rows):
    """按各列内容显示宽度占比，把 B..BC（列 2..55）动态分成 n 段。

    - 每段至少 1 列；合并边界取整列，不改任何列宽
    - 需求 = 该列表头+数据中最大显示宽度 + 2 字符余量
    - 1 列表格不使用本函数（整段 B:BC）
    """
    needs = []
    for j in range(n):
        vals = [headers[j]] + [r[j] for r in rows if j < len(r)]
        needs.append(max(disp_width(str(v or '')) for v in vals) + 2)
    total_w = seg_width(ws, 2, 55)
    total_need = sum(needs)
    widths = [col_width(ws, c) for c in range(2, 56)]
    target = [total_w * (nd / total_need) for nd in needs]

    segs = []
    idx = 0  # 已消耗列数（对应列 2+idx）
    for k, t in enumerate(target):
        if k == n - 1:
            segs.append((2 + idx, 55))  # 最后一段收尾到 BC
            break
        acc = 0.0
        start = 2 + idx
        acc += widths[idx]
        idx += 1
        # 预留后续每段至少 1 列
        while idx < len(widths) - (n - k - 1) and acc < t:
            acc += widths[idx]
            idx += 1
        segs.append((start, 2 + idx - 1))
    return segs


def col_align(values):
    if not values:
        return 'left'
    maxw = max(disp_width(str(v or '')) for v in values)
    return 'center' if maxw <= 10 else 'left'


# ---------------- MD 解析 ----------------
def _split_row(r):
    """按未转义的 | 切分单元格；\\| 还原成 |"""
    r = r.strip()
    if r.startswith('|'):
        r = r[1:]
    if r.endswith('|') and not r.endswith('\\|'):
        r = r[:-1]
    cells, cur, i = [], [], 0
    while i < len(r):
        if r[i] == '\\' and i + 1 < len(r) and r[i + 1] == '|':
            cur.append('|')
            i += 2
            continue
        if r[i] == '|':
            cells.append(''.join(cur).strip())
            cur = []
        else:
            cur.append(r[i])
        i += 1
    cells.append(''.join(cur).strip())
    return cells


def _is_sep_row(r):
    return bool(re.match(r'^\s*\|?[\s:\-|]+\|?\s*$', r)) and '-' in r


def _clean_inline(text):
    """去除行内代码反引号标记（markdown 渲染语义），保留其余原文"""
    return re.sub(r'`([^`]*)`', r'\1', str(text))


def _unescape_cell(text):
    """表格单元格文本：<br> 转为单元格内换行，\\| 还原为 |"""
    return BR_RE.sub('\n', str(text)).replace('\\|', '|')


def _cell_text(text):
    """md 表格单元格文本：去反引号标记，再做 <br> / \\| 处理"""
    return _unescape_cell(_clean_inline(text))


def _png_size(path):
    """读取 PNG 像素尺寸（读 IHDR，不依赖 Pillow）"""
    with open(path, 'rb') as f:
        f.read(16)
        w, h = struct.unpack('>II', f.read(8))
    return w, h


def _render_mermaid(code, out_png, width=MERMAID_WIDTH):
    """用 mermaid-cli 渲染 mermaid 代码为 PNG。

    返回 (ok, message)：失败时 message 为 mmdc 输出的首条错误信息，供调用方告警。
    """
    if not MMDC:
        return False, 'mmdc 未找到（需要 @mermaid-js/mermaid-cli）'
    mmd = out_png + '.mmd'
    with open(mmd, 'w', encoding='utf-8') as f:
        f.write(code)
    try:
        r = subprocess.run([MMDC, '-i', mmd, '-o', out_png, '-b', 'white', '-w', str(width)],
                           capture_output=True, text=True, timeout=MMDC_TIMEOUT)
    except (OSError, subprocess.TimeoutExpired) as e:
        return False, str(e)
    finally:
        if os.path.exists(mmd):
            os.remove(mmd)
    if r.returncode != 0 or not os.path.exists(out_png):
        lines = [ln.strip() for ln in (r.stderr + '\n' + r.stdout).splitlines() if ln.strip()]
        return False, lines[0] if lines else f'mmdc 退出码 {r.returncode}'
    return True, ''


def parse_md(text):
    lines = text.split('\n')
    blocks = []
    doc_title = None
    first_level = None
    i, n = 0, len(lines)
    while i < n:
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        m = re.match(r'^(#{1,6})\s+(.*)', line)
        if m:
            level, title = len(m.group(1)), _clean_inline(m.group(2).strip())
            if first_level is None:
                first_level = level
                doc_title = title
            elif level == first_level + 1:
                blocks.append(('chapter', title))
            elif level >= first_level + 2:
                blocks.append(('section', title))
            else:  # 同级后续标题按章节处理
                blocks.append(('chapter', title))
            i += 1
            continue
        if line.lstrip().startswith('```'):
            fence = line.strip()[3:].strip().lower()
            i += 1
            code = []
            while i < n and not lines[i].strip().startswith('```'):
                code.append(lines[i])
                i += 1
            i += 1  # 跳过结束 ```
            body = '\n'.join(code).rstrip()
            blocks.append(('mermaid', body) if fence == 'mermaid' else ('code', body))
            continue
        if line.strip().startswith('|'):
            tbl = []
            while i < n and lines[i].strip().startswith('|'):
                tbl.append(lines[i])
                i += 1
            blocks.append(parse_table(tbl))
            continue
        para = []
        while i < n:
            l = lines[i].rstrip()
            if not l.strip() or l.strip().startswith('|') or l.strip().startswith('```') \
                    or re.match(r'^(#{1,6})\s+', l):
                break
            para.append(_clean_inline(l.strip()))
            i += 1
        if para:
            blocks.append(('note', '\n'.join(para)))
    return doc_title, blocks


def parse_table(rows):
    headers = [_cell_text(h) for h in _split_row(rows[0])]
    data = []
    for r in rows[1:]:
        if _is_sep_row(r):
            continue
        data.append([_cell_text(c) for c in _split_row(r)])
    maxc = max([len(headers)] + [len(r) for r in data])
    headers = headers + [''] * (maxc - len(headers))
    data = [r + [''] * (maxc - len(r)) for r in data]
    return ('table', headers, data)


def normalize_table(headers, rows):
    """>3 列时并入第 3 列（表头/内容用 ' / ' 连接），返回 (headers, rows)"""
    n = len(headers)
    if n <= 3:
        return headers, rows
    h3 = ' / '.join(headers[2:])
    headers = headers[:2] + [h3]
    out = []
    for r in rows:
        v3 = ' / '.join(str(x) for x in r[2:] if str(x) != '')
        out.append(r[:2] + [v3])
    return headers, out


# ---------------- 渲染 ----------------
def build(template_path, out_path, blocks, doc_title=None, sheet_name=None):
    wb = load_workbook(template_path)
    ws = wb.active
    if sheet_name is None:
        sheet_name = _make_sheet_name(doc_title or ws['A1'].value)
    ws.title = sheet_name

    # 清理模板内容区（第 7 行起），保留前 6 行与列宽
    if ws.max_row > 6:
        ws.delete_rows(7, ws.max_row - 6)
    # delete_rows 不清理行高：去掉模板遗留的第 7 行之后的行高，否则空行会被撑高
    for r in [r for r in ws.row_dimensions if r > 6]:
        del ws.row_dimensions[r]
    ws.merged_cells.ranges = [mr for mr in ws.merged_cells.ranges if mr.max_row <= 6]
    for row in ws.iter_rows(min_row=1, max_row=6):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith('=#REF!'):
                cell.value = None
    ws.sheet_state = 'visible'
    for s in list(wb.sheetnames):
        if s != sheet_name:
            del wb[s]
    wb.active = 0

    if doc_title:
        ws['A1'] = doc_title

    header_fill = PatternFill('solid', fgColor='FF666666')
    section_fill = PatternFill('solid', fgColor='FFD9D9D9')
    code_fill = PatternFill('solid', fgColor='FFF2F2F2')
    img_dir = tempfile.mkdtemp(prefix='mermaid_', dir=str(default_temp_dir()))

    def set_merged(r, c1, c2, value, fill, font, align, bordered):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            if fill is not None:
                cell.fill = fill
            cell.font = font
            cell.alignment = align
            cell.border = (Border(left=LINE_THIN, right=LINE_THIN,
                                  top=LINE_THIN, bottom=LINE_THIN) if bordered else Border())
        if c2 > c1:
            ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
        ws.cell(row=r, column=c1).value = value

    row = 7
    for block in blocks:
        kind = block[0]
        if kind == 'chapter':
            set_merged(row, FULL[0], FULL[1], block[1], header_fill,
                       Font(name=FONT_JA, size=11, bold=True, color='FFFFFFFF'),
                       Alignment(horizontal='left', vertical='center'), bordered=False)
            ws.row_dimensions[row].height = 24
        elif kind == 'section':
            set_merged(row, FULL[0], FULL[1], block[1], section_fill,
                       Font(name=FONT_JA, size=10, bold=True, color='FF000000'),
                       Alignment(horizontal='left', vertical='center'), bordered=False)
            ws.row_dimensions[row].height = 22
        elif kind == 'note':
            set_merged(row, FULL[0], FULL[1], block[1], None,
                       Font(name=FONT_JA, size=10, color='FF000000'),
                       Alignment(horizontal='left', vertical='center', wrap_text=True),
                       bordered=False)
            ws.row_dimensions[row].height = max(18, est_lines(block[1], *FULL, ws) * NOTE_LINE_PT + NOTE_PAD)
        elif kind == 'code':
            set_merged(row, FULL[0], FULL[1], block[1], code_fill,
                       Font(name=FONT_MONO, size=10, color='FF333333'),
                       Alignment(horizontal='left', vertical='center', wrap_text=True),
                       bordered=False)
            ws.row_dimensions[row].height = est_lines(block[1], *FULL, ws) * CODE_LINE_PT + CODE_PAD
        elif kind == 'mermaid':
            png = os.path.join(img_dir, f'mermaid_{row}.png')
            ok, msg = _render_mermaid(block[1], png)
            if ok:
                w, h = _png_size(png)
                scale = min(1.0, MERMAID_MAX_H / h)
                img = XLImage(png)
                img.width = int(w * scale)
                img.height = int(h * scale)
                ws.add_image(img, f'B{row}')
                ws.row_dimensions[row].height = img.height * 0.75 + 4
            else:  # mmdc 不可用或渲染失败 → 降级为代码文本，并告警
                print(f'warning: 第 {row} 行 mermaid 渲染失败，已降级为代码文本: {msg}', file=sys.stderr)
                set_merged(row, FULL[0], FULL[1], block[1], code_fill,
                           Font(name=FONT_MONO, size=10, color='FF333333'),
                           Alignment(horizontal='left', vertical='center', wrap_text=True),
                           bordered=False)
                ws.row_dimensions[row].height = est_lines(block[1], *FULL, ws) * CODE_LINE_PT + CODE_PAD
        elif kind == 'table':
            headers, rows = normalize_table(block[1], block[2])
            n = len(headers)
            segs = SEG1 if n == 1 else dynamic_segments(ws, n, headers, rows)
            for j, h in enumerate(headers):
                c1, c2 = segs[j]
                set_merged(row, c1, c2, h, header_fill,
                           Font(name=FONT_JA, size=10, bold=True, color='FFFFFFFF'),
                           Alignment(horizontal='center', vertical='center', wrap_text=True),
                           bordered=True)
            ws.row_dimensions[row].height = 22
            row += 1
            for r in rows:
                for j, v in enumerate(r):
                    c1, c2 = segs[j]
                    set_merged(row, c1, c2, v, None,
                               Font(name=FONT_JA, size=9, color='FF000000'),
                               Alignment(horizontal=col_align([rr[j] for rr in rows]),
                                         vertical='center', wrap_text=True),
                               bordered=True)
                    ws.row_dimensions[row].height = max(
                        ws.row_dimensions[row].height or 0,
                        est_lines(v, c1, c2, ws) * TABLE_LINE_PT + TABLE_PAD)
                row += 1
        else:
            raise ValueError(f'未知的 block 类型: {kind!r}')
        row += 1  # 块间空行

    wb.save(out_path)
    shutil.rmtree(img_dir, ignore_errors=True)
    return ws.max_row


def _make_sheet_name(title):
    name = (title or DEFAULT_SHEET).strip()
    name = re.sub(r'[\\/*?:\[\]]', '', name).replace('：', '_').replace('/', '_')
    return (name or DEFAULT_SHEET)[:31]


def find_template(cli_tpl):
    """模板查找顺序：--template 显式 > 本 Skill assets/design_template.xlsx（唯一权威）"""
    if cli_tpl:
        return Path(cli_tpl)
    return ASSET_TPL


def default_out_dir():
    d = Path.cwd() / 'output'
    d.mkdir(parents=True, exist_ok=True)
    return d


def default_temp_dir():
    d = Path.cwd() / 'temp'
    d.mkdir(parents=True, exist_ok=True)
    return d


def load_blocks_json(path):
    data = json.loads(Path(path).read_text(encoding='utf-8'))
    out = []
    for i, item in enumerate(data):
        kind = item[0] if item else None
        if kind not in BLOCK_KINDS:
            raise ValueError(f'blocks[{i}]: 未知的 block 类型 {kind!r}（允许: {", ".join(BLOCK_KINDS)}）')
        if kind == 'table':
            if len(item) != 3 or not isinstance(item[1], list) or not isinstance(item[2], list):
                raise ValueError(f'blocks[{i}]: table 需要 ["table", [表头...], [[行...], ...]]')
            out.append((kind, [_unescape_cell(h) for h in item[1]],
                        [[_unescape_cell(c) for c in r] for r in item[2]]))
        else:
            if len(item) != 2 or not isinstance(item[1], str):
                raise ValueError(f'blocks[{i}]: {kind} 需要 ["{kind}", "文本"]')
            out.append((kind, item[1]))
    return out


def main():
    ap = argparse.ArgumentParser(description='MD → 設計書 Excel（design_template 版式）')
    ap.add_argument('--input', help='输入 Markdown 文件路径')
    ap.add_argument('--blocks', help='blocks JSON 文件路径（与 --input 二选一）')
    ap.add_argument('--output', help='输出 xlsx 路径（默认：./output/<源文件名>.xlsx）')
    ap.add_argument('--template', help='模板 xlsx 路径（默认：本 Skill assets/design_template.xlsx）')
    ap.add_argument('--sheet-name', help='sheet 名（默认取文档标题）')
    ap.add_argument('--title', help='A1 文档标题（默认：MD 模式取首个 # 标题，blocks 模式保留模板标题）')
    args = ap.parse_args()

    if not args.input and not args.blocks:
        ap.error('必须提供 --input 或 --blocks')
    tpl = find_template(args.template)
    if not tpl.exists():
        sys.exit(f'模板不存在: {tpl}')

    out_dir = default_out_dir()
    if args.blocks:
        blocks = load_blocks_json(args.blocks)
        out = Path(args.output) if args.output else out_dir / f'{Path(args.blocks).stem}.xlsx'
        doc_title = args.title  # None 则保留模板 A1 标题
    else:
        src = Path(args.input)
        md_title, blocks = parse_md(src.read_text(encoding='utf-8'))
        out = Path(args.output) if args.output else out_dir / f'{src.stem}.xlsx'
        doc_title = args.title or md_title
        if doc_title is None and blocks and blocks[0][0] == 'chapter':
            doc_title = blocks[0][1]

    rows = build(str(tpl), str(out), blocks, doc_title=doc_title, sheet_name=args.sheet_name)
    print(f'saved: {out} ({rows} rows, sheet={Path(out).stem})')


if __name__ == '__main__':
    main()
