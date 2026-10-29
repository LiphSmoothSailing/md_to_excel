# -*- coding: utf-8 -*-
"""设计书 Excel 版式校验：回读生成的 xlsx，对照 design_template.xlsx。

检查项：
  1. 单 sheet 且 visible；给了 --sheet-name 时校验 sheet 名
  2. 列宽与模板完全一致（显式列宽 + defaultColWidth）
  3. 前 6 行：合并区与模板相同；值与模板相同（A1 标题、=#REF! 清理两处除外）；行高相同
  4. 第 7 行起没有残留的模板行高（无内容、无合并、无图片的行不得有行高）
  5. 合并区互不重叠
  6. 报告 mermaid 图片数量（信息项，不判定）

用法：
    python verify_design_excel.py --output <目标.xlsx> [--template <模板.xlsx>] [--sheet-name <名>]

退出码 0 = 全部通过；1 = 有失败项。
"""
import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import md_to_design_excel as D  # noqa: E402

HEAD_ROWS = 6

fails = []


def check(name, cond, extra=''):
    print(('PASS  ' if cond else 'FAIL  ') + name + ('' if cond else f'   <- {extra}'))
    if not cond:
        fails.append(name)


def col_widths(ws):
    return {k: v.width for k, v in ws.column_dimensions.items() if v.width}


def used_rows(ws):
    """有内容、合并或图片的行号集合"""
    rows = set()
    for row in ws.iter_rows(min_row=HEAD_ROWS + 1):
        for c in row:
            if c.value is not None:
                rows.add(c.row)
    for mr in ws.merged_cells.ranges:
        rows.update(range(mr.min_row, mr.max_row + 1))
    for img in getattr(ws, '_images', []):
        rows.add(img.anchor._from.row + 1)
    return rows


def main():
    ap = argparse.ArgumentParser(description='设计书 Excel 版式校验')
    ap.add_argument('--output', required=True, help='待校验的 xlsx')
    ap.add_argument('--template', default=str(D.ASSET_TPL), help='模板 xlsx（默认 Skill assets）')
    ap.add_argument('--sheet-name', help='期望的 sheet 名')
    args = ap.parse_args()

    tpl = load_workbook(args.template).active
    out_wb = load_workbook(args.output)
    ws = out_wb.active

    check('单 sheet', len(out_wb.sheetnames) == 1, out_wb.sheetnames)
    check('sheet 可见', ws.sheet_state == 'visible', ws.sheet_state)
    if args.sheet_name:
        check(f'sheet 名 = {args.sheet_name}', ws.title == args.sheet_name, ws.title)

    # 列宽
    check('defaultColWidth 与模板一致',
          ws.sheet_format.defaultColWidth == tpl.sheet_format.defaultColWidth,
          f'{ws.sheet_format.defaultColWidth} vs {tpl.sheet_format.defaultColWidth}')
    tw, ow = col_widths(tpl), col_widths(ws)
    diff = {k: (tw.get(k), ow.get(k)) for k in set(tw) | set(ow) if tw.get(k) != ow.get(k)}
    check('显式列宽与模板完全一致', not diff, diff)

    # 前 6 行：合并区
    tpl_merges = {str(m) for m in tpl.merged_cells.ranges if m.max_row <= HEAD_ROWS}
    out_merges = {str(m) for m in ws.merged_cells.ranges if m.max_row <= HEAD_ROWS}
    check(f'前 6 行合并区与模板相同（{len(tpl_merges)} 个）', tpl_merges == out_merges,
          f'缺 {sorted(tpl_merges - out_merges)} 多 {sorted(out_merges - tpl_merges)}')

    # 前 6 行：值与行高
    bad_vals = []
    for r in range(1, HEAD_ROWS + 1):
        for c in range(1, tpl.max_column + 1):
            if (r, c) == (1, 1):
                continue                      # A1 标题允许改写
            tv = tpl.cell(r, c).value
            if isinstance(tv, str) and tv.startswith('=#REF!'):
                tv = None                     # 断链公式按设计清空
            if tv != ws.cell(r, c).value:
                bad_vals.append(ws.cell(r, c).coordinate)
    check('前 6 行值与模板一致（A1 与 #REF! 除外）', not bad_vals, bad_vals[:10])
    bad_h = [r for r in range(1, HEAD_ROWS + 1)
             if tpl.row_dimensions[r].height != ws.row_dimensions[r].height]
    check('前 6 行行高与模板一致', not bad_h, bad_h)

    # 第 7 行起无残留行高
    used = used_rows(ws)
    stale = sorted(r for r, d in ws.row_dimensions.items()
                   if r > HEAD_ROWS and d.height and r not in used)
    check('第 7 行起无残留模板行高', not stale, f'行 {stale[:10]}')

    # 合并区重叠
    rngs = list(ws.merged_cells.ranges)
    overlap = None
    for i, a in enumerate(rngs):
        for b in rngs[i + 1:]:
            if not (a.max_row < b.min_row or b.max_row < a.min_row
                    or a.max_col < b.min_col or b.max_col < a.min_col):
                overlap = (str(a), str(b))
    check('合并区互不重叠', overlap is None, overlap)

    print(f'info  图片 {len(getattr(ws, "_images", []))} 张，最大行 {ws.max_row}')
    print()
    print('FAILED:', fails if fails else 'none')
    sys.exit(1 if fails else 0)


if __name__ == '__main__':
    main()
