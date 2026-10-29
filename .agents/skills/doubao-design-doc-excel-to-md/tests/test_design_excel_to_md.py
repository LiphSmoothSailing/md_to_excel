# -*- coding: utf-8 -*-
"""design_excel_to_md 自测：md → 设计书 xlsx（doubao-design-doc-excel）→ md 往返校验。

运行：
    python .agents\\skills\\doubao-design-doc-excel-to-md\\tests\\test_design_excel_to_md.py
退出码 0 = 全部通过。中间产物写在系统临时目录下。
"""
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve()
SKILL_DIR = HERE.parents[1]                                   # doubao-design-doc-excel-to-md
FWD_DIR = HERE.parents[2] / 'doubao-design-doc-excel' / 'scripts'
sys.path.insert(0, str(SKILL_DIR / 'scripts'))
sys.path.insert(0, str(FWD_DIR))
import design_excel_to_md as R  # noqa: E402
import md_to_design_excel as F  # noqa: E402

from openpyxl import load_workbook  # noqa: E402
from openpyxl.drawing.image import Image as XLImage  # noqa: E402
from PIL import Image as PILImage  # noqa: E402

fails = []


def check(name, cond, extra=''):
    print(('PASS  ' if cond else 'FAIL  ') + name + ('' if cond else f'   <- {extra}'))
    if not cond:
        fails.append(name)


SRC = """# 测试设计书

## 1. 概要

第一段。
第二段。

### 1.1 项目

| 项目 | 说明 | 备注 |
| --- | --- | --- |
| a | x<br>y | p \\| q |

```json
{"a": 1}
```

## 2. 结论

结束。
"""

TMP = Path(tempfile.mkdtemp(prefix='d2m_'))
xlsx = TMP / 'design.xlsx'
title, blocks = F.parse_md(SRC)
F.build(str(F.ASSET_TPL), str(xlsx), blocks, doc_title=title)

ws = load_workbook(xlsx).active
lines, images = R.convert(ws, 'design_images')
text = '\n'.join(lines)

check('A1 标题 → # 标题', lines[0] == '# 测试设计书', lines[0])
check('章节 → ##', '## 1. 概要' in lines and '## 2. 结论' in lines)
check('子节 → ###', '### 1.1 项目' in lines)
check('正文多行保留换行（两段分别成行）', '第一段。' in lines and '第二段。' in lines)
check('表头行', '| 项目 | 说明 | 备注 |' in lines)
check('表格单元格换行 → <br>', '| a | x<br>y | p \\| q |' in lines,
      [ln for ln in lines if ln.startswith('| a')])
check('管道符转义 \\| 保留', 'p \\| q' in text)
check('代码块带 json 语言标记', '```json' in lines and '{"a": 1}' in lines)
check('代码块围栏闭合', lines[lines.index('```json') + 2] == '```')
check('模板头部保留为 HTML 注释', '<!-- 版式头部（模板信息）' in lines)
check('无图片时没有导出项', images == [])

# 往返：再转回 xlsx 应得到相同的块结构（标题/章节/子节/表格/代码）
xlsx2 = TMP / 'design2.xlsx'
title2, blocks2 = F.parse_md('\n'.join(lines))
F.build(str(F.ASSET_TPL), str(xlsx2), blocks2, doc_title=title2)
ws2 = load_workbook(xlsx2).active
lines2, _ = R.convert(ws2, 'design_images')
core = lambda ls: [ln for ln in ls if ln.startswith(('#', '| ', '```', '{')) and not ln.startswith('<!--')]
check('往返：标题/章节/子节/表格/代码结构一致', core(lines) == core(lines2),
      f'{core(lines)[:6]} vs {core(lines2)[:6]}')

# 图片：内嵌 PNG 导出为文件并以相对路径引用
img_ws = load_workbook(xlsx).active
png = TMP / 'pixel.png'
PILImage.new('RGB', (8, 8), 'white').save(png)
img_ws.add_image(XLImage(str(png)), 'B20')
img_xlsx = TMP / 'img.xlsx'
img_ws.parent.save(img_xlsx)
img_ws2 = load_workbook(img_xlsx).active
lines3, images3 = R.convert(img_ws2, 'img_dir')
check('图片导出 1 张', len(images3) == 1, len(images3))
check('图片以相对路径引用', any(ln.startswith('![流程图 1](img_dir/mermaid_1.png)') for ln in lines3))

# 对真实样例（如存在）做冒烟：能转换、结构数与源文件标题数一致
real = Path(sys.argv[1]) if len(sys.argv) > 1 else None
if real and real.is_file():
    rb = F.parse_md(real.read_text(encoding='utf-8'))
    rx = TMP / 'real.xlsx'
    F.build(str(F.ASSET_TPL), str(rx), rb[1], doc_title=rb[0])
    rl, _ = R.convert(load_workbook(rx).active, 'real_images')
    src_heads = [ln for ln in real.read_text(encoding='utf-8').split('\n') if ln.startswith('## ')]
    got_heads = [ln for ln in rl if ln.startswith('## ')]
    check(f'真实文件（{real.name}）：章节数一致', len(src_heads) == len(got_heads),
          f'{len(src_heads)} vs {len(got_heads)}')
else:
    print('SKIP  真实文件冒烟（未提供 md 路径）')

print()
print('FAILED:', fails if fails else 'none')
sys.exit(1 if fails else 0)
