# -*- coding: utf-8 -*-
"""設計書 Excel → Markdown（与 doubao-design-doc-excel 互逆）。

按 design_template 版式的样式识别各块（B:BC 为内容区）：
  - 标题（A1）          → `# 标题`
  - 章节                 B:BC 整段合并 + 深灰底 FF666666 + 白字粗体、无边框  → `## 标题`
  - 子节                 B:BC 整段合并 + 浅灰底 FFD9D9D9                     → `### 标题`
  - 表格                 B 列有细左边框的连续行；首行为表头（深灰底）          → GFM 管道表格
  - 代码                 B:BC 整段合并 + 浅灰底 FFF2F2F2 + Consolas          → 代码围栏
  - 正文                 B:BC 整段合并、无底色、无边框                       → 段落
  - 图片                 工作表内嵌图片（mermaid 渲染结果）                   → 导出 PNG 并引用

第 1–6 行的模板头部（サービス名、作成者等）作为 HTML 注释保留，不进入正文。
表格单元格：换行 → `<br>`，`|` → `\\|`。

限制（Excel 中本来就不存在的信息无法还原）：
  - mermaid 源码：只能导出图片，无法还原为 mermaid 代码块
  - 行内 `code` 反引号、**粗体** 在设计书转换时已被去除，这里不会恢复
  - 表格对齐标记统一输出为 `---`；超过 3 列并入第 3 列的表格不会拆回去
  - 第 4 级及以下标题在转换时已降为子节，这里输出为 `###`

用法：
    python design_excel_to_md.py --input <源.xlsx> [--output <目标.md>] [--force]

图片缺省导出到 `<目标文件名>_images/` 目录（与 .md 同级）。
"""
import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

FIRST_COL, LAST_COL = 2, 55          # B..BC 内容区
HEAD_ROWS = 6

HEADER_FILL = 'FF666666'
SECTION_FILL = 'FFD9D9D9'
CODE_FILL = 'FFF2F2F2'
CODE_FONT = 'Consolas'


def fill_rgb(cell):
    f = cell.fill
    if f is None or f.fill_type != 'solid':
        return None
    return f.fgColor.rgb


def has_thin_left(cell):
    b = cell.border
    return bool(b and b.left and b.left.style == 'thin')


def merge_maps(ws):
    """anchors: {(row, col): max_col} 合并区左上角；covered: 合并区内非左上角的格"""
    anchors, covered = {}, set()
    for mr in ws.merged_cells.ranges:
        for r in range(mr.min_row, mr.max_row + 1):
            for c in range(mr.min_col, mr.max_col + 1):
                if (r, c) == (mr.min_row, mr.min_col):
                    anchors[(r, c)] = mr.max_col
                else:
                    covered.add((r, c))
    return anchors, covered


def row_segments(anchors, covered, r):
    """按合并区把一行切成列段 [(起始列, 结束列)]，与生成时的表格列段一一对应"""
    segs, c = [], FIRST_COL
    while c <= LAST_COL:
        if (r, c) in covered:
            c += 1
            continue
        end = anchors.get((r, c), c)
        segs.append((c, end))
        c = end + 1
    return segs


def text_of(ws, r, c):
    v = ws.cell(r, c).value
    return '' if v is None else str(v)


def classify(ws, r, image_rows):
    if r in image_rows:
        return 'image'
    b = ws.cell(r, FIRST_COL)
    if has_thin_left(b):
        return 'table'
    if b.value in (None, ''):
        return 'blank'
    fill = fill_rgb(b)
    if fill == HEADER_FILL:
        return 'chapter'
    if fill == SECTION_FILL:
        return 'section'
    if fill == CODE_FILL or b.font.name == CODE_FONT:
        return 'code'
    return 'note'


def esc_cell(text):
    return str(text).replace('|', '\\|').replace('\n', '<br>')


def table_lines(rows):
    n = len(rows[0])
    out = ['| ' + ' | '.join(esc_cell(h) for h in rows[0]) + ' |',
           '| ' + ' | '.join(['---'] * n) + ' |']
    for r in rows[1:]:
        cells = (r + [''] * n)[:n]
        out.append('| ' + ' | '.join(esc_cell(c) for c in cells) + ' |')
    return out


def code_lines(text):
    lang = ''
    stripped = text.strip()
    if stripped[:1] in '{[':
        try:
            json.loads(stripped)
            lang = 'json'
        except ValueError:
            pass
    fence = '````' if '```' in text else '```'
    return [f'{fence}{lang}', text, fence]


def header_comment(ws):
    """第 1–6 行模板头部（A1 除外）的非空值，每行一条，用 HTML 注释保留"""
    lines = []
    for r in range(1, HEAD_ROWS + 1):
        vals = [text_of(ws, r, c) for c in range(1, ws.max_column + 1)
                if not (r == 1 and c == 1) and ws.cell(r, c).value not in (None, '')]
        if vals:
            lines.append(' / '.join(vals))
    return lines


def convert(ws, images_dir_name):
    """工作表 → (markdown 行列表, 图片列表 [(文件名, bytes)])"""
    anchors, covered = merge_maps(ws)
    images = sorted(ws._images, key=lambda img: img.anchor._from.row)
    image_rows = {img.anchor._from.row + 1: i for i, img in enumerate(images)}
    image_files, lines = [], []

    title = ws['A1'].value
    if title:
        lines += [f'# {title}', '']
    comment = header_comment(ws)
    if comment:
        lines += ['<!-- 版式头部（模板信息）'] + comment + ['-->', '']

    last_row = max([ws.max_row] + list(image_rows))   # 图片可能锚定在内容行之后
    r = HEAD_ROWS + 1
    while r <= last_row:
        kind = classify(ws, r, image_rows)
        if kind == 'blank':
            r += 1
            continue
        if kind == 'table':
            segs = row_segments(anchors, covered, r)
            rows = []
            while r <= ws.max_row and has_thin_left(ws.cell(r, FIRST_COL)):
                rows.append([text_of(ws, r, c1) for c1, _ in segs])
                r += 1
            lines += table_lines(rows) + ['']
            continue
        if kind == 'image':
            idx = image_rows[r]
            name = f'mermaid_{idx + 1}.png'
            image_files.append((name, images[idx]._data()))
            lines += [f'![流程图 {idx + 1}]({images_dir_name}/{name})', '']
        else:
            text = text_of(ws, r, FIRST_COL)
            if kind == 'chapter':
                lines += [f'## {text}', '']
            elif kind == 'section':
                lines += [f'### {text}', '']
            elif kind == 'code':
                lines += code_lines(text) + ['']
            else:
                lines += text.split('\n') + ['']
        r += 1

    while lines and lines[-1] == '':
        lines.pop()
    return lines, image_files


def main():
    ap = argparse.ArgumentParser(description='設計書 Excel → Markdown（与 doubao-design-doc-excel 互逆）')
    ap.add_argument('--input', required=True, help='输入 xlsx（design_template 版式）')
    ap.add_argument('--output', help='输出 md（缺省 ./output/<源文件名>.md）')
    ap.add_argument('--force', action='store_true', help='输出已存在时允许覆盖（缺省拒绝）')
    args = ap.parse_args()

    src = Path(args.input)
    if not src.is_file():
        sys.exit(f'输入文件不存在: {src}')
    out = Path(args.output) if args.output else Path.cwd() / 'output' / (src.stem + '.md')
    if out.exists() and not args.force:
        sys.exit(f'输出文件已存在，未覆盖：{out}  （确认覆盖加 --force）')
    out.parent.mkdir(parents=True, exist_ok=True)

    ws = load_workbook(src).active
    img_dir_name = f'{out.stem}_images'
    lines, image_files = convert(ws, img_dir_name)
    out.write_text('\n'.join(lines) + '\n', encoding='utf-8')
    if image_files:
        img_dir = out.parent / img_dir_name
        img_dir.mkdir(parents=True, exist_ok=True)
        for name, data in image_files:
            (img_dir / name).write_bytes(data)

    print(f'input  : {src}')
    print(f'saved  : {out}  (lines={len(lines)})')
    print(f'images : {len(image_files)} 张' + (f'，导出到 {out.parent / img_dir_name}' if image_files else ''))


if __name__ == '__main__':
    main()
