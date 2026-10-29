# -*- coding: utf-8 -*-
"""Excel → Markdown（与 md-to-excel-plain 互逆）：

  - 有细边框的行 → GFM 管道表格（表头 + 分隔行 + 数据行）
  - A 列无边框的行 → 原文照搬为 Markdown 文本行（标题 #、列表、加粗 ** 等都保留）
  - 空行保留

用法：
    python excel_to_md_plain.py --input <源.xlsx> [--output <目标.md>]

参数：
    --input PATH    输入 xlsx（必需）
    --output PATH   输出 md（缺省 ./output/<源文件名>.md）
    --force         输出已存在时允许覆盖
"""
import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.utils import range_boundaries


def build_merge_map(ws):
    """掃描所有合併區，返回 (covered, anchors)。

    covered: set[(r, c)] —— 落在合併區內但不是左上角錨點的格（讀表時應跳過）。
    anchors: dict[(r, c)] -> (min_col, max_col) —— 該格是某段橫向合併的起點。
    """
    covered = set()
    anchors = {}
    for mr in ws.merged_cells.ranges:
        min_col, min_row, max_col, max_row = range_boundaries(str(mr))
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                if r == min_row and c == min_col:
                    anchors[(r, c)] = (min_col, max_col)
                else:
                    covered.add((r, c))
    return covered, anchors


def cell_text(val):
    """单元格值 → 字符串（None → 空串）"""
    if val is None:
        return ""
    return str(val)


def rich_to_md(val, skip_bold=False):
    """单元格值 → Markdown 文本；粗体 run 包成 **...**。

    - None → 空串
    - 普通字符串 → 原样（里面的 ** 本来就是字面量，保留）
    - CellRichText → 逐 run 拼接；TextBlock.font.b=True 的 run 包 **，相邻粗体 run 合并
    - skip_bold=True 时不包（用于标题行 / 表头这类结构性加粗，避免 # 标题被包成 **#...**）
    """
    if val is None:
        return ""
    if not isinstance(val, CellRichText):
        return str(val)
    parts = []
    prev_bold = False
    for blk in val:
        if isinstance(blk, TextBlock):
            is_bold = bool(getattr(blk.font, "b", False)) and not skip_bold
            text = blk.text or ""
        else:
            is_bold = False
            text = str(blk)
        if not text:
            continue
        if is_bold and not prev_bold:
            parts.append("**")
        if not is_bold and prev_bold:
            parts.append("**")
        parts.append(text)
        prev_bold = is_bold
    if prev_bold:
        parts.append("**")
    return "".join(parts)


def is_heading_text(val):
    """是否为 Markdown 标题行（# 开头）——这类行的加粗是结构性的，不包 **"""
    s = str(val or "").lstrip()
    return s.startswith("#")


def escape_pipe(s):
    """把单元格里的字面 | 转义为 \\|（GFM 表格语法要求），单元格内换行还原为 <br>"""
    return s.replace("|", "\\|").replace("\n", "<br>")


def col_align_from_cells(header_cell, data_cell):
    """根据表头和数据单元格的水平对齐，反推 GFM 分隔行标记。

    md-to-excel-plain 的对应关系：
      源 ---  (无对齐)  → header=center, data=left
      源 :---               → header=left,   data=left
      源 :---:              → header=center, data=center
      源 ---:               → header=right,  data=right
    """
    ha = header_cell.alignment.horizontal if header_cell and header_cell.alignment else None
    da = data_cell.alignment.horizontal if data_cell and data_cell.alignment else None
    if ha == "left" and da == "left":
        return ":---"
    if ha == "center" and da == "center":
        return ":---:"
    if ha == "right" and da == "right":
        return "---:"
    # header=center + data=left（或其它）→ 无对齐
    return "---"


def row_is_table(ws, r, max_col):
    """该行是否为表格行（A 列有细左边框）"""
    c = ws.cell(row=r, column=1)
    return c.border is not None and c.border.left is not None and c.border.left.style == "thin"


def row_text(ws, r):
    """非表格行：取 A 列文本；标题行的结构性加粗不包 **"""
    v = ws.cell(row=r, column=1).value
    return rich_to_md(v, skip_bold=is_heading_text(v))


def table_row_cells(ws, r, ncol, skip_bold=False):
    """表格行：取前 ncol 列文本（無合併場景用）"""
    return [rich_to_md(ws.cell(row=r, column=c).value, skip_bold=skip_bold)
            for c in range(1, ncol + 1)]


def table_row_cells_at(ws, r, col_starts, skip_bold=False):
    """表格行：在指定邏輯列位置上取值（合併格已由 col_starts 去重）"""
    return [rich_to_md(ws.cell(row=r, column=c).value, skip_bold=skip_bold)
            for c in col_starts]


def table_to_md(header_cells, data_rows, aligns):
    """把一张表渲染成 GFM 管道表格文本块"""
    lines = []
    lines.append("| " + " | ".join(escape_pipe(c) for c in header_cells) + " |")
    lines.append("| " + " | ".join(aligns) + " |")
    for row in data_rows:
        lines.append("| " + " | ".join(escape_pipe(c) for c in row) + " |")
    return lines


def convert(ws):
    """遍历工作表，返回 Markdown 文本行列表"""
    covered, anchors = build_merge_map(ws)
    max_row = ws.max_row
    max_col = ws.max_column
    out = []
    r = 1
    while r <= max_row:
        if not row_is_table(ws, r, max_col):
            out.append(row_text(ws, r))
            r += 1
            continue
        # 收集连续的表格行
        block_start = r
        while r <= max_row and row_is_table(ws, r, max_col):
            r += 1
        block_end = r - 1
        # 邏輯列：表頭行上的合併錨點列 + 有值的獨格列；跳過合併區內非錨點格
        col_starts = []
        for c in range(1, max_col + 1):
            if (block_start, c) in covered:
                continue
            val = ws.cell(row=block_start, column=c).value
            is_anchor = (block_start, c) in anchors
            if (val not in (None, "")) or is_anchor:
                col_starts.append(c)
        ncol = len(col_starts)
        if ncol == 0:
            ncol = 1
            col_starts = [1]
        header = table_row_cells_at(ws, block_start, col_starts, skip_bold=True)
        data = [table_row_cells_at(ws, rr, col_starts, skip_bold=False)
                for rr in range(block_start + 1, block_end + 1)]
        # 反推對齊：用表頭行和第一條數據行在各邏輯列位置上的對齊
        aligns = []
        for c in col_starts:
            hc = ws.cell(row=block_start, column=c)
            dc = ws.cell(row=block_start + 1, column=c) if data else None
            aligns.append(col_align_from_cells(hc, dc))
        out.extend(table_to_md(header, data, aligns))
    return out


def main():
    ap = argparse.ArgumentParser(description="Excel → Markdown（表格转管道表，其余原文照搬）")
    ap.add_argument("--input", required=True, help="输入 xlsx 路径")
    ap.add_argument("--output", help="输出 md 路径（缺省 ./output/<源文件名>.md）")
    ap.add_argument("--force", action="store_true", help="输出已存在时允许覆盖")
    args = ap.parse_args()

    src = Path(args.input)
    if not src.is_file():
        sys.exit(f"输入文件不存在: {src}")
    out = Path(args.output) if args.output else Path.cwd() / "output" / (src.stem + ".md")
    if out.exists() and not args.force:
        sys.exit(f"输出文件已存在，未覆盖：{out}  （确认覆盖加 --force）")
    out.parent.mkdir(parents=True, exist_ok=True)

    wb = load_workbook(src, rich_text=True)
    ws = wb.active
    lines = convert(ws)
    text = "\n".join(lines)
    if not text.endswith("\n"):
        text += "\n"
    out.write_text(text, encoding="utf-8")
    n_tables = sum(1 for i, ln in enumerate(lines)
                   if ln.startswith("| ") and i + 1 < len(lines)
                   and set(lines[i + 1].replace("|", "").replace(" ", "").replace(":", "")) == {"-"})
    print(f"input  : {src}")
    print(f"rows   : {ws.max_row}")
    print(f"tables : {n_tables} 张")
    print(f"saved  : {out}  (lines={len(lines)})")


if __name__ == "__main__":
    main()
