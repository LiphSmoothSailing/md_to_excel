# -*- coding: utf-8 -*-
"""excel-to-md-plain 自测：造一个小 xlsx → 转 md → 校验输出。"""
import sys
from pathlib import Path

HERE = Path(__file__).resolve()
SKILL_DIR = HERE.parents[1]
sys.path.insert(0, str(SKILL_DIR / "scripts"))
import excel_to_md_plain as E

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

THIN = Side(style="thin", color="FF808080")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

ROOT = HERE.parents[4] if len(HERE.parents) > 4 else SKILL_DIR
TEMP = ROOT / "temp"
TEMP.mkdir(parents=True, exist_ok=True)

fails = []

def check(name, cond, extra=""):
    print(("PASS  " if cond else "FAIL  ") + name + ("" if cond else f"   <- {extra}"))
    if not cond:
        fails.append(name)

# 构造一个小 xlsx
wb = Workbook()
ws = wb.active
ws.title = "test"
# 文本行
ws.cell(row=1, column=1, value="# 見出し")
ws.cell(row=2, column=1, value="")
ws.cell(row=3, column=1, value="段落 **bold** と `code`。")
ws.cell(row=4, column=1, value="")
# 表格
ws.cell(row=5, column=1, value="名前").border = BORDER
ws.cell(row=5, column=2, value="点数").border = BORDER
ws.cell(row=5, column=1).font = Font(bold=True)
ws.cell(row=5, column=2).font = Font(bold=True)
ws.cell(row=5, column=1).alignment = Alignment(horizontal="center")
ws.cell(row=5, column=2).alignment = Alignment(horizontal="center")
ws.cell(row=6, column=1, value="田中").border = BORDER
ws.cell(row=6, column=2, value="90").border = BORDER
ws.cell(row=6, column=1).alignment = Alignment(horizontal="left")
ws.cell(row=6, column=2).alignment = Alignment(horizontal="left")
# 空行 + 文本
ws.cell(row=7, column=1, value="")
ws.cell(row=8, column=1, value="- リスト項目")

lines = E.convert(ws)
check("行数 = 9（6 文本行 + 表 3 行）", len(lines) == 9, f"got {len(lines)}: {lines}")
check("# 見出し 原样保留", lines[0] == "# 見出し", lines[0])
check("空行保留", lines[1] == "", repr(lines[1]))
check("段落含 ** 原样", lines[2] == "段落 **bold** と `code`。", lines[2])
check("表头行", lines[4] == "| 名前 | 点数 |", lines[4])
check("分隔行", lines[5] == "| --- | --- |", lines[5])
check("数据行", lines[6] == "| 田中 | 90 |", lines[6])
check("列表保留", lines[8] == "- リスト項目", lines[8])

# 管道转义
ws2 = Workbook().active
ws2.cell(row=1, column=1, value="a|b").border = BORDER
ws2.cell(row=1, column=2, value="正常").border = BORDER
ws2.cell(row=2, column=1, value="c").border = BORDER
ws2.cell(row=2, column=2, value="d").border = BORDER
lines2 = E.convert(ws2)
check("管道转义 \\|", lines2[0] == "| a\\|b | 正常 |", lines2[0])

# 居中对齐往返：:---: 不能变成 ---:
ws3 = Workbook().active
for r in (1, 2):
    for c, v in enumerate(["标题", "值"] if r == 1 else ["a", "1"], start=1):
        cell = ws3.cell(row=r, column=c, value=v)
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center")
lines3 = E.convert(ws3)
check("居中对齐 :---: 往返不变成 ---:", lines3[1] == "| :---: | :---: |", lines3[1])

# 反斜杠原样保留（不能翻倍）
ws4 = Workbook().active
ws4.cell(row=1, column=1, value="p\\q").border = BORDER
ws4.cell(row=1, column=2, value="x").border = BORDER
ws4.cell(row=2, column=1, value="a").border = BORDER
ws4.cell(row=2, column=2, value="b").border = BORDER
lines4 = E.convert(ws4)
check("反斜杠原样保留", lines4[0] == "| p\\q | x |", lines4[0])

# 单元格内换行还原为 <br>
ws5 = Workbook().active
ws5.cell(row=1, column=1, value="列").border = BORDER
ws5.cell(row=2, column=1, value="a\nb").border = BORDER
lines5 = E.convert(ws5)
check("单元格内换行还原为 <br>", lines5[2] == "| a<br>b |", lines5[2])

print()
print("FAILED:", fails if fails else "none")
sys.exit(1 if fails else 0)
